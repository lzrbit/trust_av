"""Exploratory controlled-factor distributions and repeated-participant inference.

Portable input/output paths; never edits input data or existing results.
The v8 family is 195 paired clip comparisons and 12 joint-factor Wald tests.
These families do not replace the v5 75-contrast bootstrap family.
"""
from pathlib import Path
from datetime import datetime, timezone
from itertools import combinations
import argparse, hashlib, json
import numpy as np
import pandas as pd
from scipy import stats, linalg, special
from openpyxl import load_workbook

from settings import WORKSPACE as ROOT, DATA_DIR, REPO, relative_label
import os
BLOCKS = {'LC':'B1','SVM':'B2','HB':'B3','MB':'B4'}
SHEETS = {'LC':'MAL','SVM':'AMB','HB':'HB','MB':'MB'}
SPECS = {'LC': {'lateral_behaviour':['1m/s','3m/s','Fragmented','Abortion'], 'acc_style':['cautious','mild','aggressive'], 'design_distance':[5,15]}}
for _s in ['SVM','HB','MB']:
    SPECS[_s] = {'design_speed':[80,100,120], 'design_braking':[-2,-5,-8], 'design_distance':[5,15,25]}
FIELD = {'design_speed':'design_speed_kmh','design_braking':'design_braking_m_s2','design_distance':'design_distance_m','lateral_behaviour':'lateral_behaviour','acc_style':'acc_style'}
UNIT = {'design_speed':'km/h','design_braking':'m/s2','design_distance':'m','lateral_behaviour':'category','acc_style':'category'}
DISPLAY = {'1m/s':'Normal, 1 m/s','3m/s':'Normal, 3 m/s','Fragmented':'Fragmented','Abortion':'Aborted','cautious':'Cautious','mild':'Mild','aggressive':'Aggressive'}

def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def dump(path, obj):
    Path(path).write_text(json.dumps(obj, ensure_ascii=False, indent=2, allow_nan=False)+'\n')

def holm(p):
    p=np.asarray(p,float); order=np.argsort(p,kind='stable'); out=np.empty(len(p))
    out[order]=np.minimum(1.,np.maximum.accumulate(p[order]*(len(p)-np.arange(len(p)))))
    return out

def log_t_twosided(t, df):
    value=float(np.log(2.)+stats.t.logsf(abs(t),df))
    if np.isfinite(value): return min(0.,value)
    a=df/2.; x=df/(df+t*t)
    # Regularised incomplete-beta identity, evaluated on the log scale.
    value=a*np.log(x)-np.log(a)-special.betaln(a,.5)+np.log(special.hyp2f1(a,.5,a+1,x))
    assert np.isfinite(value)
    return min(0.,float(value))

def log_chi2_tail(stat, df):
    # All planned df are even; Gamma integer shape gives an exact finite sum.
    assert df%2==0
    k=np.arange(df//2); x=stat/2.
    return float(-x+special.logsumexp(k*np.log(x)-special.gammaln(k+1)))

def holm_log(logp):
    logp=np.asarray(logp,float); order=np.argsort(logp,kind='stable'); out=np.empty(len(logp))
    out[order]=np.minimum(0.,np.maximum.accumulate(logp[order]+np.log(len(logp)-np.arange(len(logp)))))
    return out

def stars(p):
    return '***' if p<.001 else '**' if p<.01 else '*' if p<.05 else 'ns'

def display_level(value):
    return DISPLAY.get(value, str(value))

def extract_config(xlsx, design, dest):
    """Preserve exact source cells/headers for canonical design sheets, not raw traces."""
    wb=load_workbook(xlsx,read_only=True,data_only=False)
    inventory=[]; cells=[]; configs=[]
    for ws in wb:
        canonical=ws.title in SHEETS.values()
        inventory.append({'sheet':ws.title,'rows':ws.max_row,'columns':ws.max_column,'canonical_design':canonical,'role': 'canonical event configuration' if canonical else ('unlabelled numerical array; not interpreted as event configuration' if ws.title=='MAL_New' else 'alternative/duplicate LC design worksheet; not used for current assignments')})
        if not canonical: continue
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cells.append({'sheet':ws.title,'cell':cell.coordinate,'row':cell.row,'column':cell.column,'data_type':cell.data_type,'raw_value':str(cell.value)})
    for scenario,sheet in SHEETS.items():
        ws=wb[sheet]; index={str(row[0]): (n,row) for n,row in enumerate(ws.iter_rows(values_only=True),1) if row[0] is not None}
        sd=design[design.display_scenario.eq(scenario)].sort_values('questionnaire_event_id')
        assert len(sd)==(24 if scenario=='LC' else 27)
        for r in sd.itertuples():
            rownum,raw=index[r.source_event_name]
            assert rownum==r.source_excel_row
            vals=[r.lateral_behaviour,r.acc_style,r.design_distance_m] if scenario=='LC' else [r.design_speed_kmh,r.design_braking_m_s2,r.design_distance_m]
            assert all(str(v)==str(raw[j+1]) or (isinstance(raw[j+1],(int,float)) and float(v)==raw[j+1]) for j,v in enumerate(vals))
            configs.append({'scenario':r.scenario,'display_scenario':scenario,'event_key':r.event_key,'source_sheet':sheet,'source_row':rownum,'source_event_name':r.source_event_name,'factor1_source_header':ws.cell(1,2).value,'factor1_value':raw[1],'factor2_source_header':ws.cell(1,3).value,'factor2_value':raw[2],'factor3_source_header':ws.cell(1,4).value,'factor3_value':raw[3],'lc_distance_conflict':bool(r.lc_distance_conflict),'stored_ax_n_conflict':bool(r.stored_ax_n_conflict)})
    wb.close()
    pd.DataFrame(cells).to_csv(dest/'tables/parameter_configuration_source_cells.csv',index=False)
    pd.DataFrame(configs).to_csv(dest/'tables/parameter_configurations.csv',index=False)
    design.to_csv(dest/'tables/parameter_event_design.csv',index=False)
    dump(dest/'reports/parameter_workbook_inventory.json',{'source_sha256':sha(xlsx),'sheets':inventory,'canonical_source_cells':len(cells),'event_configurations':len(configs),'preservation':'Source cells retain original sheet, Excel coordinates, field spelling and raw units. Historical keep/delete notes are retained in source-cell CSV but do not change current clip inclusion. Alternative sheets and unlabelled MAL_New array are inventoried, not used as canonical assignments.'})

def within_system(cell, ids, spec):
    """Each observed person's event weights sum to one within the clip."""
    d=cell[cell.risk.notna()].copy(); pos=pd.Index(ids).get_indexer(d.participant_id); assert (pos>=0).all()
    n=np.bincount(pos,minlength=len(ids)); y=d.risk.to_numpy(float); meta=[]; columns=[]
    for factor,levels in spec.items():
        for level in levels[1:]:
            meta.append((factor,levels[0],level)); columns.append(d[FIELD[factor]].eq(level).to_numpy(float))
    x=np.column_stack(columns); q=x.shape[1]; sx=np.zeros((len(ids),q)); np.add.at(sx,pos,x)
    sy=np.bincount(pos,weights=y,minlength=len(ids)); xc=x-sx[pos]/n[pos,None]; yc=y-sy[pos]/n[pos]
    ai=np.zeros((len(ids),q,q)); bi=np.zeros((len(ids),q)); ci=np.bincount(pos,weights=yc*yc/n[pos],minlength=len(ids))
    np.add.at(ai,pos,np.einsum('ni,nj->nij',xc,xc)/n[pos,None,None]); np.add.at(bi,pos,xc*yc[:,None]/n[pos,None])
    a=ai.sum(0); b=bi.sum(0); assert np.linalg.matrix_rank(a)==q
    beta=np.linalg.solve(a,b); score=bi-np.einsum('nij,j->ni',ai,beta)
    assert np.max(np.abs(score.sum(0)))<1e-8
    return {'n':n,'A':a,'b':b,'c':ci.sum(),'beta':beta,'score':score,'meta':meta,'n_ratings':len(d),'xc':xc,'yc':yc,'pos':pos,'data':d}

def frozen_plan(paths, dest):
    plan={'created_utc':datetime.now(timezone.utc).isoformat(),'status':'run-local reproduction specification; not a new preregistration','selection':'HB initial braking distance (5/15/25 m) chosen for unambiguous experimental field, complete temporal profile and coverage; not selected by new p values or maximum effect. All factors retained in SI.','primary_scope':'Exploratory post-review analysis; fixed current stimulus library; assigned design levels; no causal or population-stimulus generalisation.','inputs':{k:{'path':str(v),'sha256':sha(v)} for k,v in paths.items()},'boxplot':'Within participant, average all positive (>0) event ratings at each scenario/factor level/clip. Box quantiles use numpy linear interpolation, whiskers actual extrema within 1.5IQR. No fliers removed.','paired_family':{'size':195,'test':'two-sided one-sample t on participant-wise paired level means; comparison minus reference; t interval is pointwise','correction':'Holm across all 195 tests, including 15 in main HB-distance panel','independence_unit':'participant; levels marginal over other available design factors, not fully matched on those factors'},'omnibus_family':{'size':12,'null':'all nonreference level coefficients at every clip for the factor equal zero','models':'Separate weighted participant fixed-intercept additive categorical-factor regression within each scenario/clip, exactly preserving the 21 existing FE estimates. Each participant total weight 1 within clip.','covariance':'Stack the participant scores across clips, preserving off-diagonal blocks. Sandwich covariance with G/(G-1) correction, no additional observation-level df multiplier. G is participants with any positive rating in that scenario.','reference_distribution':'asymptotic chi-squared, df = number of tested coefficients (LC lateral18/ACC12/distance6; all others10). Fail explicitly on rank deficiency.','correction':'Holm12','effect_size':'Weighted within-person partial R2 = (sum weighted within SSE_reduced - SSE_full)/SSE_reduced. Reduced model drops that factor in each clip, retaining all other categorical factors and person intercepts. This is descriptive conditional fit, not predictive R2 or old ANOVA eta-squared.'},'known_sensitivities':'Existing LC19–24 distance-exclusion results remain separate and unchanged. HB7–9 stored ax_n inconsistency is not used as grouping; assigned braking agrees with velocity derivative. No re-estimation of historical 75-pair bootstrap family.'}
    path=dest/'reports/parameter_analysis_plan.json'
    if path.exists():
        old=json.loads(path.read_text()); assert old['inputs']==plan['inputs'], 'Existing frozen plan has changed input hashes'; return old
    dump(path,plan); return plan

def run_analysis(project_root, results_dir=None, data_dir=None, design_csv=None, parameter_xlsx=None):
    root=Path(project_root).resolve(); dest=Path(results_dir or root/'outputs/parameter_inference').resolve()
    for name in ['tables','figures','reports']: (dest/name).mkdir(parents=True,exist_ok=True)
    dd=Path(data_dir or DATA_DIR); old=root/'outputs/parameters/tables'
    paths={'windows':dd/'windows.parquet','participants':dd/'participants.parquet','design':Path(design_csv or REPO/'data/event_design.csv'),'configuration_export':REPO/'data/parameter_configurations.csv','previous_profiles':old/'controlled_parameter_clip_profiles.csv','previous_adjusted':old/'controlled_parameter_adjusted_sensitivity.csv'}
    if parameter_xlsx: paths['parameter_workbook']=Path(parameter_xlsx)
    plan=frozen_plan(paths,dest); hashes={k:sha(p) for k,p in paths.items()}
    people=pd.read_parquet(paths['participants']); ids=np.sort(people.participant_id.to_numpy()); assert len(ids)==len(set(ids))==2164
    w=pd.read_parquet(paths['windows']); assert len(w)==181776 and w.risk.notna().sum()==179966 and w.risk_raw.eq(0).sum()==1810
    assert w.loc[w.risk.notna(),'risk'].between(1,10).all()
    design=pd.read_csv(paths['design']); assert len(design)==105 and design.event_key.nunique()==105
    if parameter_xlsx:
        extract_config(Path(parameter_xlsx),design,dest)
    else:
        # The included table is a validated, stimulus-only export of the workbook.
        # This branch reproduces the declared design; it does not re-verify raw cells.
        configuration=pd.read_csv(paths['configuration_export'])
        assert len(configuration)==105 and configuration.event_key.is_unique
        check=configuration.merge(design,on='event_key',validate='one_to_one',suffixes=('_configuration','_design'))
        assert len(check)==105
        for row in check.itertuples():
            expected=[row.lateral_behaviour,row.acc_style,row.design_distance_m] if row.display_scenario_design=='LC' else [row.design_speed_kmh,row.design_braking_m_s2,row.design_distance_m]
            actual=[row.factor1_value,row.factor2_value,row.factor3_value]
            for a,b in zip(actual,expected):
                try:assert float(a)==float(b)
                except ValueError:assert str(a)==str(b)
        configuration.to_csv(dest/'tables/parameter_configurations.csv',index=False)
        design.to_csv(dest/'tables/parameter_event_design.csv',index=False)
        dump(dest/'reports/configuration_export_validation.json',{'status':'pass','events':105,'source_sha256':sha(paths['configuration_export']),'original_workbook_reverified':False,'scope':'Included stimulus-only configuration export agrees with all 105 design rows. Supply --parameter-xlsx to verify original workbook cells.'})
    w=w.merge(design,on=['event_key','scenario'],how='left',validate='many_to_one',suffixes=('','_design')); assert w.source_event_name.notna().all()
    boxes=[]; pairs=[]; omnibus=[]; coeff=[]; cover=[]; covariances={}; candidate=[]
    for code,spec in SPECS.items():
        sub=w[w.scenario.eq(BLOCKS[code])]; clips=range(1,7 if code=='LC' else 6)
        for factor,levels in spec.items():
            col=FIELD[factor]; means=sub.groupby(['participant_id',col,'window_index']).risk.mean(); matrices=[]
            for li,level in enumerate(levels):
                mat=means.xs(level,level=col).unstack('window_index').reindex(index=ids,columns=clips).to_numpy(); matrices.append(mat)
                for ci,clip in enumerate(clips):
                    a=mat[:,ci]; a=a[np.isfinite(a)]; q1,median,q3=np.quantile(a,[.25,.5,.75],method='linear'); iqr=q3-q1; inside=a[(a>=q1-1.5*iqr)&(a<=q3+1.5*iqr)]; low,high=inside.min(),inside.max(); fliers=a[(a<low)|(a>high)]
                    boxes.append(dict(scenario=BLOCKS[code],display_scenario=code,factor=factor,level=level,level_label=display_level(level),level_order=li+1,unit=UNIT[factor],clip=clip,n_participants=len(a),mean=a.mean(),sd=a.std(ddof=1),q1=q1,median=median,q3=q3,whisker_low=low,whisker_high=high,n_fliers=len(fliers),fliers_json=json.dumps(fliers.tolist(),separators=(',',':')),row_key=f'{code}|{factor}|{level}|{clip}'))
            for a,b in combinations(range(len(levels)),2):
                delta=matrices[b]-matrices[a]
                for ci,clip in enumerate(clips):
                    d=delta[:,ci]; d=d[np.isfinite(d)]; n=len(d); mean=d.mean(); sd=d.std(ddof=1); se=sd/np.sqrt(n); assert n>1 and se>0
                    t=mean/se; p=float(2*stats.t.sf(abs(t),n-1)); critical=stats.t.ppf(.975,n-1)
                    pairs.append(dict(scenario=BLOCKS[code],display_scenario=code,factor=factor,clip=clip,reference_level=levels[a],comparison_level=levels[b],reference_order=a+1,comparison_order=b+1,n_paired_participants=n,mean_difference=mean,paired_sd=sd,se=se,ci_low=mean-critical*se,ci_high=mean+critical*se,t_statistic=t,df=n-1,p_raw=p,log10_p_raw=log_t_twosided(t,n-1)/np.log(10.),p_raw_underflow=p==0.,contrast_definition='comparison minus reference; same-person factor-level means; other factors marginal',row_key=f'{code}|{factor}|{levels[b]}-{levels[a]}|{clip}'))
        systems=[within_system(sub[sub.window_index.eq(clip)],ids,spec) for clip in clips]
        q=len(systems[0]['beta']); G=int(np.any(np.stack([s['n']>0 for s in systems]),axis=0).sum()); assert G>1
        beta=np.concatenate([s['beta'] for s in systems]); scores=np.concatenate([s['score'] for s in systems],axis=1); bread=linalg.block_diag(*[np.linalg.inv(s['A']) for s in systems]); covariance=(G/(G-1))*bread@(scores.T@scores)@bread.T; covariance=(covariance+covariance.T)/2
        assert np.linalg.eigvalsh(covariance).min()>-1e-12
        covariances[code]=covariance
        sse_full=sum(s['c']-2*s['beta']@s['b']+s['beta']@s['A']@s['beta'] for s in systems)
        for ki,(clip,s) in enumerate(zip(clips,systems)):
            cover.append(dict(display_scenario=code,scenario=BLOCKS[code],clip=clip,n_participants=int((s['n']>0).sum()),n_with_two_or_more=int((s['n']>=2).sum()),n_valid_ratings=s['n_ratings'],n_parameters=q,design_rank=int(np.linalg.matrix_rank(s['A']))))
            for j,(factor,ref,lev) in enumerate(s['meta']):
                index=ki*q+j; se=np.sqrt(covariance[index,index]); coeff.append(dict(display_scenario=code,scenario=BLOCKS[code],clip=clip,factor=factor,reference_level=ref,comparison_level=lev,estimate=beta[index],cluster_se=se,ci_low=beta[index]-stats.norm.ppf(.975)*se,ci_high=beta[index]+stats.norm.ppf(.975)*se,cluster_n=G,covariance_column=index))
        for factor,levels in spec.items():
            local=[j for j,m in enumerate(systems[0]['meta']) if m[0]==factor]; ix=[ki*q+j for ki in range(len(systems)) for j in local]; v=covariance[np.ix_(ix,ix)]; rank=int(np.linalg.matrix_rank(v)); assert rank==len(ix), f'Non-estimable Wald: {code}/{factor}'
            b=beta[ix]; stat=float(b@np.linalg.solve(v,b)); p=float(stats.chi2.sf(stat,len(ix)))
            keep=[j for j in range(q) if j not in local]; sse_reduced=0.
            for s in systems:
                ar=s['A'][np.ix_(keep,keep)]; br=s['b'][keep]; cr=np.linalg.solve(ar,br); sse_reduced+=s['c']-2*cr@br+cr@ar@cr
            partial=(sse_reduced-sse_full)/sse_reduced; assert -1e-12<=partial<=1
            omnibus.append(dict(scenario=BLOCKS[code],display_scenario=code,factor=factor,reference_level=levels[0],levels_json=json.dumps(levels),n_clips=len(systems),n_level_coefficients=len(levels)-1,wald_chi2=stat,df=len(ix),covariance_rank=rank,cluster_n=G,n_valid_ratings=sum(s['n_ratings'] for s in systems),cluster_correction=G/(G-1),p_raw=p,log10_p_raw=log_chi2_tail(stat,len(ix))/np.log(10.),p_raw_underflow=p==0.,weighted_sse_full=sse_full,weighted_sse_reduced=sse_reduced,weighted_within_partial_r2=max(0.,partial),null='All factor coefficients at all clips are zero',row_key=f'{code}|{factor}'))
        # Independent residual-score reconstruction for one clip, no score shortcut.
        s=systems[0]; residual=s['yc']-s['xc']@s['beta']; direct=np.zeros_like(s['score']); np.add.at(direct,s['pos'],s['xc']*residual[:,None]/s['n'][s['pos'],None]); assert np.max(np.abs(direct-s['score']))<1e-12
    bdf=pd.DataFrame(boxes); pdf=pd.DataFrame(pairs); odf=pd.DataFrame(omnibus); cdf=pd.DataFrame(coeff)
    assert len(bdf)==189 and len(pdf)==195 and len(odf)==12 and len(cdf)==126
    pdf['log10_p_holm_195']=holm_log(pdf.log10_p_raw.to_numpy()*np.log(10.))/np.log(10.); pdf['p_holm_195']=np.power(10.,pdf.log10_p_holm_195); pdf['significance']=pdf.p_holm_195.map(stars); pdf['family_size']=195
    odf['log10_p_holm_12']=holm_log(odf.log10_p_raw.to_numpy()*np.log(10.))/np.log(10.); odf['p_holm_12']=np.power(10.,odf.log10_p_holm_12); odf['family_size']=12
    previous=pd.read_csv(paths['previous_profiles']); previous=previous[previous.analysis_variant.eq('assigned_design_primary')].copy()
    for frame in [previous,bdf]: frame['level']=frame.level.astype(str).str.replace(r'\.0$','',regex=True)
    check=previous.merge(bdf,on=['display_scenario','factor','level','clip'],validate='one_to_one',suffixes=('_old','_new')); assert len(check)==189
    mean_error=float(np.max(np.abs(check.mean_old-check.mean_new))); assert mean_error<1e-12 and check.n_participants_old.eq(check.n_participants_new).all()
    previous_a=pd.read_csv(paths['previous_adjusted']); previous_a=previous_a[previous_a.analysis_variant.eq('assigned_design_primary')].copy()
    for frame in [previous_a,cdf]:
        for col in ['reference_level','comparison_level']: frame[col]=frame[col].astype(str).str.replace(r'\.0$','',regex=True)
    check_a=previous_a.merge(cdf,on=['display_scenario','factor','clip','reference_level','comparison_level'],validate='one_to_one',suffixes=('_old','_new')); assert len(check_a)==126
    beta_error=float(np.max(np.abs(check_a.estimate_old-check_a.estimate_new))); assert beta_error<1e-12
    for (code,factor),g in bdf.groupby(['display_scenario','factor'],sort=False):
        candidate.append({'display_scenario':code,'factor':factor,'n_min':int(g.n_participants.min()),'n_max':int(g.n_participants.max()),'mean_min':g['mean'].min(),'mean_max':g['mean'].max(),'selected_main_panel':code=='HB' and factor=='design_distance','selection_basis':'Design clarity, full clip pattern and coverage; no new inferential p values used in selection'})
    outputs={'parameter_boxstats.csv':bdf,'parameter_paired_tests.csv':pdf,'parameter_omnibus_tests.csv':odf,'parameter_adjusted_coefficients.csv':cdf,'parameter_model_coverage.csv':pd.DataFrame(cover),'parameter_candidate_inventory.csv':pd.DataFrame(candidate)}
    for name,frame in outputs.items(): frame.to_csv(dest/'tables'/name,index=False)
    bdf[bdf.display_scenario.eq('HB')&bdf.factor.eq('design_distance')].to_csv(dest/'tables/figure_1h_hb_distance_boxstats.csv',index=False)
    pdf[pdf.display_scenario.eq('HB')&pdf.factor.eq('design_distance')].to_csv(dest/'tables/figure_1h_hb_distance_paired_tests.csv',index=False)
    np.savez_compressed(dest/'reports/parameter_joint_covariances.npz',**covariances)
    assert all(sha(p)==hashes[k] for k,p in paths.items())
    report={'status':'pass','created_utc':datetime.now(timezone.utc).isoformat(),'plan_sha256':sha(dest/'reports/parameter_analysis_plan.json'),'source_hashes':{k:{'path':str(p),'sha256':hashes[k]} for k,p in paths.items()},'source_files_unchanged':True,'participant_n':len(ids),'positive_ratings':int(w.risk.notna().sum()),'zero_no_operation_ratings':int(w.risk_raw.eq(0).sum()),'counts':{k:len(v) for k,v in outputs.items()},'previous_estimates_reproduction':{'profile_rows':189,'maximum_mean_difference':mean_error,'all_profile_sample_counts_equal':True,'adjusted_coefficients':126,'maximum_coefficient_difference':beta_error},'paired_star_counts':pdf.significance.value_counts().to_dict(),'HB_distance_star_counts':pdf.loc[pdf.display_scenario.eq('HB')&pdf.factor.eq('design_distance'),'significance'].value_counts().to_dict(),'omnibus_underflow_count':int(odf.p_raw_underflow.sum()),'omnibus_display_rule':'Display p<0.001 in the compact table. Numerical p=0 is underflow, not exact zero. Finite log10_p_raw/log10_p_holm preserve the actual tail scale, using the integer-gamma finite sum for Wald and incomplete-beta identity for the t tail. Every underflow here is smaller than 1e-300.','script_sha256':sha(Path(__file__)),'output_hashes':{str(p.relative_to(dest)):sha(p) for p in [*sorted((dest/'tables').glob('*.csv')),dest/'reports/parameter_joint_covariances.npz']}}
    dump(dest/'reports/parameter_analysis_validation.json',report)
    print(json.dumps({k:report[k] for k in ['status','counts','previous_estimates_reproduction','paired_star_counts','HB_distance_star_counts']},indent=2))
    return report

if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__); p.add_argument('--project-root',type=Path,default=ROOT); p.add_argument('--results-dir',type=Path); p.add_argument('--data-dir',type=Path); p.add_argument('--design-csv',type=Path); p.add_argument('--parameter-xlsx',type=Path,default=Path(os.environ['TRUST_AV_PARAMETER_XLSX']) if os.environ.get('TRUST_AV_PARAMETER_XLSX') else None); a=p.parse_args()
    run_analysis(a.project_root,a.results_dir,a.data_dir,a.design_csv,a.parameter_xlsx)
